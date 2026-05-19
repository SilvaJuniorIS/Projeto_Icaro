from __future__ import annotations

import re
from html import escape
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.config import OUTPUT_DIR
from src.db import export_snapshot, now_iso
from src.checklist import gerar_checklist
from src.comparability import avaliar_comparabilidade


def _safe_name(value: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9_-]+", "_", value.strip())[:60].strip("_")
    return slug or "processo"


def _auto_width(ws) -> None:
    for column_cells in ws.columns:
        width = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(width + 2, 12), 60)


def _style_header(ws, row: int = 1) -> None:
    fill = PatternFill("solid", fgColor="1F4E5F")
    for cell in ws[row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")


def _fmt_money(value: Any) -> str:
    if value is None:
        return "-"
    return f"R$ {float(value):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _texto(value: Any) -> str:
    text = str(value or "").strip()
    return text or "-"


def gerar_relatorio_markdown(processo_id: int, output_dir: Path | str = OUTPUT_DIR) -> Path | None:
    snapshot = export_snapshot(processo_id)
    if not snapshot:
        return None

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    processo: dict[str, Any] = snapshot["processo"]
    resumo: dict[str, Any] = snapshot["resumo"]
    resumo_atas: dict[str, Any] = snapshot["resumo_atas"]
    itens: list[dict[str, Any]] = snapshot.get("itens", [])
    resumo_itens: list[dict[str, Any]] = snapshot.get("resumo_itens", [])
    fontes: list[dict[str, Any]] = snapshot["fontes"]
    atas: list[dict[str, Any]] = snapshot["atas"]
    checklist = gerar_checklist(snapshot)
    comparabilidade = avaliar_comparabilidade(snapshot)

    linhas = [
        f"# Relatorio de pesquisa de mercado - {_texto(processo.get('titulo'))}",
        "",
        f"Gerado em: {snapshot['gerado_em']}",
        "",
        "## 1. Identificacao do processo",
        "",
        f"- Objeto: {_texto(processo.get('descricao_objeto'))}",
        f"- Categoria: {_texto(processo.get('categoria'))}",
        f"- Unidade: {_texto(processo.get('unidade'))}",
        f"- Quantidade estimada: {_texto(processo.get('quantidade'))}",
        f"- Local de execucao/entrega: {_texto(processo.get('local_execucao'))}",
        f"- Responsavel: {_texto(processo.get('responsavel'))}",
        f"- Status: {_texto(processo.get('status'))}",
        "",
        "## 2. Cesta de itens pesquisados",
        "",
        "| Codigo | Item | Unidade | Quantidade | Termo de busca | Fontes aproveitadas | Mediana |",
        "| --- | --- | --- | ---: | --- | ---: | ---: |",
    ]
    if resumo_itens:
        for item_resumo in resumo_itens:
            item = item_resumo["item"]
            item_res = item_resumo["resumo"]
            linhas.append(
                f"| {_texto(item.get('codigo'))} | {_texto(item.get('descricao'))} | {_texto(item.get('unidade'))} | {_texto(item.get('quantidade'))} | {_texto(item.get('termo_busca'))} | {item_res.get('fontes_aproveitadas', 0)} | {_fmt_money(item_res.get('preco_estimado_mediana'))} |"
            )
    else:
        linhas.append("| - | Nenhum item cadastrado | - | - | - | - | - |")

    linhas.extend(
        [
            "",
            "## 3. Pesquisa de precos consolidada",
            "",
            f"- Fontes registradas: {resumo.get('fontes_total', 0)}",
            f"- Fontes aproveitadas: {resumo.get('fontes_aproveitadas', 0)}",
            f"- Fontes descartadas: {resumo.get('fontes_descartadas', 0)}",
            f"- Menor valor: {_fmt_money(resumo.get('menor'))}",
            f"- Maior valor: {_fmt_money(resumo.get('maior'))}",
            f"- Media: {_fmt_money(resumo.get('media'))}",
            f"- Mediana: {_fmt_money(resumo.get('mediana'))}",
            f"- Preco estimado sugerido pela mediana: {_fmt_money(resumo.get('preco_estimado_mediana'))}",
            f"- Outliers identificados pelo criterio IQR: {resumo.get('outliers', 0)}",
            f"- Fontes com comparabilidade alta: {sum(1 for item in comparabilidade['itens'] if item['classificacao'] == 'alta')}",
            "",
            "### Fontes consultadas",
            "",
            "| Uso | Tipo | Item | Valor unitario | Orgao | Fornecedor | Justificativa/observacoes |",
            "| --- | --- | --- | ---: | --- | --- | --- |",
        ]
    )

    if fontes:
        for fonte in fontes:
            uso = "Aproveitada" if fonte.get("aproveitada") else "Descartada"
            observacao = fonte.get("justificativa_descarte") or fonte.get("observacoes")
            linhas.append(
                "| "
                + " | ".join(
                    [
                        uso,
                        _texto(fonte.get("fonte_tipo")),
                        _texto(fonte.get("descricao_item")),
                        _fmt_money(fonte.get("valor_unitario")),
                        _texto(fonte.get("orgao")),
                        _texto(fonte.get("fornecedor")),
                        _texto(observacao),
                    ]
                )
                + " |"
            )
    else:
        linhas.append("| - | - | Nenhuma fonte registrada | - | - | - | - |")

    fontes_com_identificacao = [
        fonte
        for fonte in fontes
        if fonte.get("descricao_item") and (fonte.get("orgao") or fonte.get("fornecedor"))
    ]
    linhas.extend(
        [
            "",
            "### Objetos localizados com identificacao da fonte",
            "",
            "| Objeto localizado | Banco/fonte | Orgao/entidade | Fornecedor | URL |",
            "| --- | --- | --- | --- | --- |",
        ]
    )
    if fontes_com_identificacao:
        for fonte in fontes_com_identificacao:
            linhas.append(
                "| "
                + " | ".join(
                    [
                        _texto(fonte.get("descricao_item")),
                        _texto(fonte.get("fonte_tipo")),
                        _texto(fonte.get("orgao")),
                        _texto(fonte.get("fornecedor") or "Nao informado na fonte"),
                        _texto(fonte.get("url")),
                    ]
                )
                + " |"
            )
    else:
        linhas.append("| Nenhum objeto com orgao ou fornecedor identificado | - | - | - | - |")

    linhas.extend(
        [
            "",
            "## 4. Atas para possivel adesao",
            "",
            f"- Atas registradas: {resumo_atas.get('atas_total', 0)}",
            f"- Atas vigentes ou sem data final informada: {resumo_atas.get('atas_vigentes', 0)}",
            f"- Atas potencialmente aderentes: {resumo_atas.get('atas_potencialmente_aderentes', 0)}",
            "",
            "| Aderencia | Ata | Orgao gerenciador | Fornecedor | Valor unitario | Vigencia | Observacoes |",
            "| --- | --- | --- | --- | ---: | --- | --- |",
        ]
    )

    if atas:
        for ata in atas:
            vigencia = f"{_texto(ata.get('vigencia_inicio'))} a {_texto(ata.get('vigencia_fim'))}"
            linhas.append(
                "| "
                + " | ".join(
                    [
                        _texto(ata.get("aderencia")),
                        _texto(ata.get("numero")),
                        _texto(ata.get("orgao_gerenciador")),
                        _texto(ata.get("fornecedor")),
                        _fmt_money(ata.get("valor_unitario")),
                        vigencia,
                        _texto(ata.get("observacoes")),
                    ]
                )
                + " |"
            )
    else:
        linhas.append("| - | Nenhuma ata registrada | - | - | - | - | - |")

    linhas.extend(
        [
            "",
            "## 5. Checklist administrativo inicial",
            "",
        ]
    )
    for item in checklist:
        linhas.append(f"- [{'x' if item['ok'] else ' '}] {item['titulo']}: {item['detalhe']}")

    linhas.extend(
        [
            "",
            "## 6. Comparabilidade",
            "",
            "| Fonte | Classificacao | Score | Similaridade | Desvio da mediana |",
            "| --- | --- | ---: | ---: | ---: |",
        ]
    )
    if comparabilidade["itens"]:
        for item in comparabilidade["itens"]:
            desvio = "-" if item["desvio_percentual_mediana"] is None else f"{item['desvio_percentual_mediana']}%"
            linhas.append(
                f"| {_texto(item['descricao_item'])} | {item['classificacao']} | {item['score']} | {item['similaridade_objeto']} | {desvio} |"
            )
    else:
        linhas.append("| Nenhuma fonte registrada | - | - | - | - |")

    linhas.extend(
        [
            "",
            "## 7. Observacao",
            "",
            "Este relatorio apoia a instrucao administrativa e nao substitui a analise tecnica, juridica ou de controle interno.",
            "",
        ]
    )

    filename = f"icaro_relatorio_{processo_id}_{_safe_name(processo['titulo'])}_{now_iso().replace(':', '-')}.md"
    path = output_dir / filename
    path.write_text("\n".join(linhas), encoding="utf-8")
    return path


def _markdown_relatorio(processo_id: int, output_dir: Path | str = OUTPUT_DIR) -> str | None:
    path = gerar_relatorio_markdown(processo_id, output_dir)
    if not path:
        return None
    return path.read_text(encoding="utf-8")


def gerar_relatorio_html(processo_id: int, output_dir: Path | str = OUTPUT_DIR) -> Path | None:
    markdown = _markdown_relatorio(processo_id, output_dir)
    if markdown is None:
        return None
    output_dir = Path(output_dir)
    processo = export_snapshot(processo_id)["processo"]
    body = []
    in_table = False
    for line in markdown.splitlines():
        if line.startswith("# "):
            if in_table:
                body.append("</table>")
                in_table = False
            body.append(f"<h1>{escape(line[2:])}</h1>")
        elif line.startswith("## "):
            if in_table:
                body.append("</table>")
                in_table = False
            body.append(f"<h2>{escape(line[3:])}</h2>")
        elif line.startswith("### "):
            if in_table:
                body.append("</table>")
                in_table = False
            body.append(f"<h3>{escape(line[4:])}</h3>")
        elif line.startswith("| ") and " --- " not in line:
            cells = [escape(cell.strip()) for cell in line.strip("|").split("|")]
            if not in_table:
                body.append("<table>")
                in_table = True
            tag = "th" if not body[-1].startswith("<tr") else "td"
            body.append("<tr>" + "".join(f"<{tag}>{cell}</{tag}>" for cell in cells) + "</tr>")
        elif line.startswith("| "):
            continue
        elif line.startswith("- "):
            if in_table:
                body.append("</table>")
                in_table = False
            body.append(f"<p>{escape(line)}</p>")
        elif line.strip():
            if in_table:
                body.append("</table>")
                in_table = False
            body.append(f"<p>{escape(line)}</p>")
    if in_table:
        body.append("</table>")
    html = f"""<!DOCTYPE html>
<html lang="pt-BR">
<head>
  <meta charset="utf-8">
  <title>Relatorio Icaro</title>
  <style>
    body {{ font-family: Arial, sans-serif; color: #172027; margin: 32px; line-height: 1.45; }}
    h1, h2, h3 {{ color: #16364d; }}
    table {{ width: 100%; border-collapse: collapse; margin: 12px 0 22px; font-size: 13px; }}
    th, td {{ border: 1px solid #d9e1e7; padding: 8px; text-align: left; vertical-align: top; }}
    th {{ background: #eef3f5; }}
    @media print {{ body {{ margin: 18mm; }} button {{ display: none; }} }}
  </style>
</head>
<body>
  <button onclick="window.print()">Imprimir / salvar em PDF</button>
  {''.join(body)}
</body>
</html>"""
    filename = f"icaro_relatorio_{processo_id}_{_safe_name(processo['titulo'])}_{now_iso().replace(':', '-')}.html"
    path = output_dir / filename
    path.write_text(html, encoding="utf-8")
    return path


def _limpar_markdown_inline(text: str) -> str:
    return (
        text.replace("**", "")
        .replace("[x]", "OK")
        .replace("[ ]", "Pendente")
        .strip()
    )


def _add_docx_table(doc: Document, linhas: list[list[str]]) -> None:
    if not linhas:
        return
    table = doc.add_table(rows=1, cols=len(linhas[0]))
    table.style = "Table Grid"
    header = table.rows[0].cells
    for idx, value in enumerate(linhas[0]):
        header[idx].text = _limpar_markdown_inline(value)
        for run in header[idx].paragraphs[0].runs:
            run.bold = True
    for linha in linhas[1:]:
        cells = table.add_row().cells
        for idx, value in enumerate(linha[: len(cells)]):
            cells[idx].text = _limpar_markdown_inline(value)
    doc.add_paragraph()


def _markdown_para_docx(markdown: str, path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    tabela_buffer: list[list[str]] = []

    def flush_table() -> None:
        nonlocal tabela_buffer
        _add_docx_table(doc, tabela_buffer)
        tabela_buffer = []

    for line in markdown.splitlines():
        stripped = line.strip()
        if not stripped:
            flush_table()
            continue
        if stripped.startswith("| "):
            if " --- " in stripped:
                continue
            cells = [cell.strip() for cell in stripped.strip("|").split("|")]
            tabela_buffer.append(cells)
            continue

        flush_table()
        if stripped.startswith("# "):
            p = doc.add_heading(_limpar_markdown_inline(stripped[2:]), level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif stripped.startswith("## "):
            doc.add_heading(_limpar_markdown_inline(stripped[3:]), level=1)
        elif stripped.startswith("### "):
            doc.add_heading(_limpar_markdown_inline(stripped[4:]), level=2)
        elif stripped.startswith("- "):
            doc.add_paragraph(_limpar_markdown_inline(stripped[2:]), style="List Bullet")
        else:
            doc.add_paragraph(_limpar_markdown_inline(stripped))

    flush_table()
    doc.save(path)


def gerar_relatorio_docx(processo_id: int, output_dir: Path | str = OUTPUT_DIR) -> Path | None:
    markdown = _markdown_relatorio(processo_id, output_dir)
    if markdown is None:
        return None
    output_dir = Path(output_dir)
    processo = export_snapshot(processo_id)["processo"]
    filename = f"icaro_relatorio_{processo_id}_{_safe_name(processo['titulo'])}_{now_iso().replace(':', '-')}.docx"
    path = output_dir / filename
    _markdown_para_docx(markdown, path)
    return path


def gerar_xlsx_processo(processo_id: int, output_dir: Path | str = OUTPUT_DIR) -> Path | None:
    snapshot = export_snapshot(processo_id)
    if not snapshot:
        return None

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    processo: dict[str, Any] = snapshot["processo"]
    resumo: dict[str, Any] = snapshot["resumo"]
    resumo_atas: dict[str, Any] = snapshot["resumo_atas"]
    itens: list[dict[str, Any]] = snapshot.get("itens", [])
    resumo_itens: list[dict[str, Any]] = snapshot.get("resumo_itens", [])
    fontes: list[dict[str, Any]] = snapshot["fontes"]
    atas: list[dict[str, Any]] = snapshot["atas"]
    checklist = gerar_checklist(snapshot)
    comparabilidade = avaliar_comparabilidade(snapshot)

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws.append(["Campo", "Valor"])
    for chave in [
        "amostras",
        "fontes_total",
        "fontes_aproveitadas",
        "fontes_descartadas",
        "menor",
        "maior",
        "media",
        "mediana",
        "preco_estimado_mediana",
        "preco_estimado_media",
        "q1",
        "q3",
        "limite_inferior",
        "limite_superior",
        "outliers",
    ]:
        ws.append([chave, resumo.get(chave)])
    for chave in ["atas_total", "atas_vigentes", "atas_potencialmente_aderentes"]:
        ws.append([chave, resumo_atas.get(chave)])
    _style_header(ws)
    _auto_width(ws)

    ws_proc = wb.create_sheet("Processo")
    ws_proc.append(["Campo", "Valor"])
    for chave, valor in processo.items():
        ws_proc.append([chave, valor])
    _style_header(ws_proc)
    _auto_width(ws_proc)

    ws_itens = wb.create_sheet("Itens")
    ws_itens.append([
        "id",
        "codigo",
        "descricao",
        "unidade",
        "quantidade",
        "categoria",
        "termo_busca",
        "fontes_aproveitadas",
        "mediana",
        "media",
        "menor",
        "maior",
    ])
    resumo_por_item = {int(item["item"]["id"]): item["resumo"] for item in resumo_itens}
    for item in itens:
        item_resumo = resumo_por_item.get(int(item["id"]), {})
        ws_itens.append([
            item.get("id"),
            item.get("codigo"),
            item.get("descricao"),
            item.get("unidade"),
            item.get("quantidade"),
            item.get("categoria"),
            item.get("termo_busca"),
            item_resumo.get("fontes_aproveitadas"),
            item_resumo.get("preco_estimado_mediana"),
            item_resumo.get("preco_estimado_media"),
            item_resumo.get("menor"),
            item_resumo.get("maior"),
        ])
    _style_header(ws_itens)
    _auto_width(ws_itens)

    ws_fontes = wb.create_sheet("Fontes")
    headers = [
        "id",
        "item_id",
        "fonte_tipo",
        "descricao_item",
        "valor_unitario",
        "quantidade",
        "orgao",
        "fornecedor",
        "uf",
        "municipio",
        "url",
        "data_referencia",
        "data_acesso",
        "aproveitada",
        "justificativa_descarte",
        "observacoes",
    ]
    ws_fontes.append(headers)
    for fonte in fontes:
        ws_fontes.append([fonte.get(header) for header in headers])
    _style_header(ws_fontes)
    _auto_width(ws_fontes)

    ws_atas = wb.create_sheet("Atas")
    ata_headers = [
        "id",
        "numero",
        "orgao_gerenciador",
        "fornecedor",
        "objeto",
        "item",
        "valor_unitario",
        "vigencia_inicio",
        "vigencia_fim",
        "quantidade_registrada",
        "quantidade_disponivel_estimativa",
        "url",
        "aderencia",
        "observacoes",
    ]
    ws_atas.append(ata_headers)
    for ata in atas:
        ws_atas.append([ata.get(header) for header in ata_headers])
    _style_header(ws_atas)
    _auto_width(ws_atas)

    ws_checklist = wb.create_sheet("Checklist")
    ws_checklist.append(["grupo", "titulo", "ok", "detalhe"])
    for item in checklist:
        ws_checklist.append([item["grupo"], item["titulo"], "sim" if item["ok"] else "nao", item["detalhe"]])
    _style_header(ws_checklist)
    _auto_width(ws_checklist)

    ws_comp = wb.create_sheet("Comparabilidade")
    ws_comp.append(["fonte_id", "item", "tipo", "valor", "similaridade", "desvio_mediana", "score", "classificacao"])
    for item in comparabilidade["itens"]:
        ws_comp.append([
            item["fonte_id"],
            item["descricao_item"],
            item["fonte_tipo"],
            item["valor_unitario"],
            item["similaridade_objeto"],
            item["desvio_percentual_mediana"],
            item["score"],
            item["classificacao"],
        ])
    _style_header(ws_comp)
    _auto_width(ws_comp)

    filename = f"icaro_processo_{processo_id}_{_safe_name(processo['titulo'])}_{now_iso().replace(':', '-')}.xlsx"
    path = output_dir / filename
    wb.save(path)
    return path
