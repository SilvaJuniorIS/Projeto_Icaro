from __future__ import annotations

import os
import tempfile
import unittest
from pathlib import Path

os.environ.setdefault("ICARO_AUTH_USER", "admin")
os.environ.setdefault("ICARO_AUTH_PASSWORD", "icaro123")
os.environ.setdefault("ICARO_AUTH_SECRET", "test-secret")

from src.db import create_fonte, create_item, create_processo, init_db
from src.reports import (
    gerar_relatorio_docx,
    gerar_relatorio_html,
    gerar_relatorio_markdown,
    gerar_xlsx_processo,
)


class ReportsTest(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.db_path = Path(self.tmp.name) / "test.sqlite3"
        self.output_dir = Path(self.tmp.name) / "output"
        init_db(self.db_path)
        self.processo_id = create_processo(
            {"titulo": "Teste Reports", "descricao_objeto": "Objeto teste"},
            self.db_path,
        )
        item_id = create_item(
            {"processo_id": self.processo_id, "descricao": "Item A", "quantidade": 10},
            self.db_path,
        )
        create_fonte(
            {
                "processo_id": self.processo_id,
                "item_id": item_id,
                "fonte_tipo": "pncp",
                "descricao_item": "Item A similar",
                "valor_unitario": 50.0,
                "aproveitada": True,
            },
            self.db_path,
        )

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def _patch_db(self):
        from unittest.mock import patch
        return patch("src.reports.export_snapshot", side_effect=self._export_snapshot)

    def _export_snapshot(self, processo_id):
        from src.db import export_snapshot
        return export_snapshot(processo_id, self.db_path)

    def test_markdown_report_generated(self) -> None:
        with self._patch_db():
            path = gerar_relatorio_markdown(self.processo_id, self.output_dir)
        self.assertIsNotNone(path)
        self.assertTrue(path.exists())
        content = path.read_text(encoding="utf-8")
        self.assertIn("Teste Reports", content)
        self.assertIn("Item A", content)

    def test_html_report_generated(self) -> None:
        with self._patch_db():
            path = gerar_relatorio_html(self.processo_id, self.output_dir)
        self.assertIsNotNone(path)
        self.assertTrue(path.exists())
        self.assertGreater(path.stat().st_size, 100)
        content = path.read_text(encoding="utf-8")
        self.assertIn("<html", content)

    def test_docx_report_generated(self) -> None:
        with self._patch_db():
            path = gerar_relatorio_docx(self.processo_id, self.output_dir)
        self.assertIsNotNone(path)
        self.assertTrue(path.exists())
        self.assertGreater(path.stat().st_size, 1000)
        from docx import Document
        doc = Document(str(path))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("Teste Reports", full_text)

    def test_xlsx_report_generated(self) -> None:
        with self._patch_db():
            path = gerar_xlsx_processo(self.processo_id, self.output_dir)
        self.assertIsNotNone(path)
        self.assertTrue(path.exists())
        from openpyxl import load_workbook
        wb = load_workbook(path)
        self.assertIn("Resumo", wb.sheetnames)
        self.assertIn("Fontes", wb.sheetnames)
        self.assertIn("Checklist", wb.sheetnames)

    def test_nonexistent_processo_returns_none(self) -> None:
        with self._patch_db():
            self.assertIsNone(gerar_relatorio_markdown(9999, self.output_dir))


if __name__ == "__main__":
    unittest.main()
