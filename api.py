from __future__ import annotations

import csv
import hmac
import os
from contextlib import asynccontextmanager
from io import BytesIO, StringIO
from typing import Any

from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from pydantic import BaseModel, Field

from src.config import BASE_DIR
from src.checklist import gerar_checklist
from src.comparability import avaliar_comparabilidade
from src.db import (
    create_ata,
    create_fonte,
    create_item,
    create_itens,
    create_fonte_from_pncp_rascunho,
    create_pncp_rascunho,
    create_processo,
    delete_ata,
    delete_fonte,
    delete_item,
    delete_pncp_rascunho,
    export_snapshot,
    get_item,
    get_pncp_rascunho,
    get_processo,
    init_db,
    list_atas,
    list_fontes,
    list_fontes_item,
    list_itens,
    list_pncp_rascunhos,
    list_processos,
    resumo_atas,
    resumo_item,
    resumo_itens,
    resumo_pesquisa,
)
from src.pncp import buscar_contratacoes, buscar_contratacoes_contextual
from src.reports import gerar_relatorio_docx, gerar_relatorio_html, gerar_relatorio_markdown, gerar_xlsx_processo
from src.review import gerar_revisao_itens


def _require_env(name: str) -> str:
    value = os.getenv(name)
    if not value:
        raise RuntimeError(
            f"Variavel de ambiente obrigatoria nao definida: {name}. "
            "Configure-a antes de iniciar a aplicacao."
        )
    return value


AUTH_USER: str = _require_env("ICARO_AUTH_USER")
AUTH_PASSWORD: str = _require_env("ICARO_AUTH_PASSWORD")
AUTH_SECRET: str = _require_env("ICARO_AUTH_SECRET")


@asynccontextmanager
async def lifespan(application: FastAPI):
    init_db()
    yield


app = FastAPI(title="Icaro", lifespan=lifespan)
app.mount("/assets", StaticFiles(directory=BASE_DIR / "assets"), name="assets")
app.mount("/icaro-docs", StaticFiles(directory=BASE_DIR / "docs"), name="icaro_docs")

AUTH_COOKIE = "icaro_session"

PUBLIC_PATHS = {
    "/",
    "/landing",
    "/icaro",
    "/atlasnex",
    "/github-page",
    "/login",
    "/auth/login",
    "/health",
}


def _assinar_usuario(usuario: str) -> str:
    assinatura = hmac.new(AUTH_SECRET.encode("utf-8"), usuario.encode("utf-8"), "sha256").hexdigest()
    return f"{usuario}:{assinatura}"


def _sessao_valida(token: str | None) -> bool:
    if not token or ":" not in token:
        return False
    usuario, assinatura = token.split(":", 1)
    assinatura_esperada = _assinar_usuario(usuario).split(":", 1)[1]
    return hmac.compare_digest(usuario, AUTH_USER) and hmac.compare_digest(assinatura, assinatura_esperada)


def _is_public_path(path: str) -> bool:
    return path in PUBLIC_PATHS or path.startswith("/assets/") or path.startswith("/icaro-docs/")


def _login_page(erro: str = "") -> str:
    erro_html = f'<div class="error">{erro}</div>' if erro else ""
    return f"""<!doctype html>
<html lang="pt-BR">
<head>
    <meta charset="utf-8">
    <meta name="viewport" content="width=device-width, initial-scale=1">
    <title>Icaro - Login</title>
    <style>
        :root {{ --bg:#dfe9ee; --ink:#172033; --muted:#586a73; --blue:#1769aa; --line:#c8d6df; --surface:#f8fbfc; }}
        * {{ box-sizing: border-box; }}
        body {{
            margin: 0; min-height: 100vh; display: grid; place-items: center; padding: 24px;
            font-family: Inter, Arial, sans-serif; color: var(--ink);
            background: linear-gradient(145deg, #d4e1e8 0%, #eef5f7 48%, #cbdbe3 100%);
        }}
        main {{
            width: min(420px, 100%); background: var(--surface); border: 1px solid var(--line);
            border-radius: 8px; padding: 28px; box-shadow: 0 18px 50px rgba(23,32,51,.16);
        }}
        h1 {{ margin: 0 0 8px; font-size: 28px; }}
        p {{ margin: 0 0 22px; color: var(--muted); line-height: 1.5; }}
        label {{ display: grid; gap: 8px; margin-bottom: 14px; font-size: 13px; font-weight: 700; color: var(--muted); }}
        input {{
            min-height: 42px; border: 1px solid var(--line); border-radius: 6px; padding: 10px 12px;
            font: inherit; color: var(--ink); background: #fff;
        }}
        button {{
            width: 100%; min-height: 42px; border: 0; border-radius: 6px; padding: 10px 14px;
            font: inherit; font-weight: 800; color: #fff; background: var(--blue); cursor: pointer;
        }}
        .links {{ display: flex; justify-content: space-between; gap: 12px; margin-top: 18px; font-size: 13px; }}
        a {{ color: var(--blue); text-decoration: none; font-weight: 700; }}
        .error {{ margin-bottom: 14px; padding: 10px 12px; border-radius: 6px; background: #fff1f1; color: #9f1d1d; font-weight: 700; }}
    </style>
</head>
<body>
    <main>
        <h1>Acessar Icaro</h1>
        <p>Entre para abrir a area de trabalho da pesquisa de mercado.</p>
        {erro_html}
        <form method="post" action="/auth/login">
            <label>Usuario<input name="usuario" autocomplete="username" required></label>
            <label>Senha<input name="senha" type="password" autocomplete="current-password" required></label>
            <button type="submit">Entrar no painel</button>
        </form>
        <div class="links">
            <a href="/">Voltar para a landing page</a>
            <a href="/icaro">Conhecer o produto</a>
        </div>
    </main>
</body>
</html>"""


@app.middleware("http")
async def exigir_login(request: Request, call_next):
    path = request.url.path
    if _is_public_path(path) or _sessao_valida(request.cookies.get(AUTH_COOKIE)):
        return await call_next(request)
    if request.method == "GET" and "text/html" in request.headers.get("accept", ""):
        return RedirectResponse(url="/login", status_code=303)
    return JSONResponse({"detail": "Autenticacao obrigatoria"}, status_code=401)


class ProcessoRequest(BaseModel):
    titulo: str
    descricao_objeto: str
    categoria: str = ""
    unidade: str = "unidade"
    quantidade: float = 1
    local_execucao: str = ""
    prazo: str = ""
    responsavel: str = ""
    status: str = "rascunho"


class FontePrecoRequest(BaseModel):
    processo_id: int
    item_id: int | None = None
    fonte_tipo: str = "pncp"
    descricao_item: str
    valor_unitario: float = Field(gt=0)
    quantidade: float = 1
    orgao: str = ""
    fornecedor: str = ""
    uf: str = ""
    municipio: str = ""
    url: str = ""
    data_referencia: str = ""
    data_acesso: str = ""
    aproveitada: bool = True
    justificativa_descarte: str = ""
    observacoes: str = ""


class AtaRequest(BaseModel):
    processo_id: int
    numero: str = ""
    orgao_gerenciador: str = ""
    fornecedor: str = ""
    objeto: str
    item: str = ""
    valor_unitario: float = Field(ge=0)
    vigencia_inicio: str = ""
    vigencia_fim: str = ""
    quantidade_registrada: float = 0
    quantidade_disponivel_estimativa: float = 0
    url: str = ""
    aderencia: str = "indefinida"
    observacoes: str = ""


class PncpBuscaRequest(BaseModel):
    termo: str
    data_inicial: str = ""
    data_final: str = ""
    uf: str = ""
    modalidade_id: str = ""
    pagina: int = 1
    tamanho_pagina: int = 10


class PncpContextoBuscaRequest(BaseModel):
    """Busca no PNCP com texto de referencia para ordenar por similaridade ao item."""

    termo: str = ""
    texto_referencia: str = ""
    item_id: int | None = None
    data_inicial: str = ""
    data_final: str = ""
    uf: str = ""
    ufs: list[str] = Field(default_factory=list)
    modalidade_id: str = ""
    tamanho_pagina: int = 14
    buscar_variantes: bool = True
    max_consultas: int = 4
    estrategia: str = "equilibrada"
    similaridade_minima: float = 0.0


class ItemPesquisaRequest(BaseModel):
    processo_id: int
    codigo: str = ""
    descricao: str
    unidade: str = ""
    quantidade: float = 1
    categoria: str = ""
    termo_busca: str = ""
    especificacao: str = ""
    status: str = "pendente"


class ItemImportInput(BaseModel):
    codigo: str = ""
    descricao: str
    unidade: str = ""
    quantidade: float = 1
    categoria: str = ""
    termo_busca: str = ""
    especificacao: str = ""
    status: str = "pendente"


class ItensImportRequest(BaseModel):
    processo_id: int
    itens: list[ItemImportInput]


class PncpRascunhoRequest(BaseModel):
    processo_id: int
    item_id: int | None = None
    numero_controle: str = ""
    objeto: str
    orgao: str = ""
    unidade: str = ""
    fornecedor: str = ""
    modalidade: str = ""
    situacao: str = ""
    data_publicacao: str = ""
    valor_estimado: float = 0
    url: str = ""
    similaridade: float = 0
    consulta: str = ""
    payload_json: dict[str, Any] | str = ""
    status: str = "rascunho"


class RascunhoFonteRequest(BaseModel):
    valor_unitario: float = Field(gt=0)
    quantidade: float = 1
    uf: str = ""
    aproveitada: bool = True
    observacoes: str = "Fonte criada a partir de rascunho PNCP."


@app.get("/", response_class=HTMLResponse)
def home() -> str:
    return (BASE_DIR / "landing.html").read_text(encoding="utf-8")


@app.get("/app", response_class=HTMLResponse)
def dashboard() -> str:
    return (BASE_DIR / "dashboard.html").read_text(encoding="utf-8")


@app.get("/login", response_class=HTMLResponse)
def login() -> str:
    return _login_page()


@app.post("/auth/login")
def auth_login(usuario: str = Form(...), senha: str = Form(...)):
    if hmac.compare_digest(usuario, AUTH_USER) and hmac.compare_digest(senha, AUTH_PASSWORD):
        response = RedirectResponse(url="/app", status_code=303)
        response.set_cookie(
            AUTH_COOKIE,
            _assinar_usuario(usuario),
            httponly=True,
            samesite="lax",
            secure=os.getenv("ICARO_COOKIE_SECURE", "0") == "1",
            max_age=60 * 60 * 12,
        )
        return response
    return HTMLResponse(_login_page("Usuario ou senha invalidos."), status_code=401)


@app.get("/auth/logout")
def auth_logout():
    response = RedirectResponse(url="/", status_code=303)
    response.delete_cookie(AUTH_COOKIE)
    return response


@app.get("/atlasnex", response_class=HTMLResponse)
def atlasnex_hub() -> str:
    return (BASE_DIR / "atlasnex.html").read_text(encoding="utf-8")


@app.get("/icaro", response_class=HTMLResponse)
def icaro_apresentacao() -> str:
    return (BASE_DIR / "icaro-index.html").read_text(encoding="utf-8")


@app.get("/landing", response_class=HTMLResponse)
def landing() -> str:
    return (BASE_DIR / "landing.html").read_text(encoding="utf-8")


@app.get("/github-page", response_class=HTMLResponse)
def github_page() -> str:
    return (BASE_DIR / "github-page.html").read_text(encoding="utf-8")


@app.get("/health")
def health() -> dict[str, str]:
    init_db()
    return {"status": "ok"}


@app.get("/processos")
def processos(limit: int = 50, offset: int = 0) -> dict[str, Any]:
    return list_processos(limit=limit, offset=offset)


@app.post("/processos")
def criar_processo(req: ProcessoRequest) -> dict[str, Any]:
    processo_id = create_processo(req.model_dump())
    return {"id": processo_id, "processo": get_processo(processo_id)}


@app.get("/processos/{processo_id}")
def obter_processo(processo_id: int) -> dict[str, Any]:
    processo = get_processo(processo_id)
    if not processo:
        raise HTTPException(status_code=404, detail="Processo nao encontrado")
    return {
        "processo": processo,
        "itens": list_itens(processo_id),
        "resumo_itens": resumo_itens(processo_id),
        "fontes": list_fontes(processo_id),
        "atas": list_atas(processo_id),
        "resumo": resumo_pesquisa(processo_id),
        "resumo_atas": resumo_atas(processo_id),
    }


@app.post("/itens")
def criar_item(req: ItemPesquisaRequest) -> dict[str, Any]:
    if not get_processo(req.processo_id):
        raise HTTPException(status_code=404, detail="Pesquisa nao encontrada")
    item_id = create_item(req.model_dump())
    return {
        "id": item_id,
        "itens": list_itens(req.processo_id),
        "resumo_itens": resumo_itens(req.processo_id),
    }


@app.post("/itens/importar")
def importar_itens(req: ItensImportRequest) -> dict[str, Any]:
    if not get_processo(req.processo_id):
        raise HTTPException(status_code=404, detail="Pesquisa nao encontrada")
    payloads = [item.model_dump() for item in req.itens]
    ids = create_itens(req.processo_id, payloads)
    return {
        "ids": ids,
        "itens": list_itens(req.processo_id),
        "resumo_itens": resumo_itens(req.processo_id),
    }


def _normalizar_cabecalho(value: Any) -> str:
    return str(value or "").strip().lower().replace("ç", "c").replace("ã", "a")


def _linha_para_item(row: dict[str, Any]) -> dict[str, Any] | None:
    descricao = str(
        row.get("descricao")
        or row.get("descrição")
        or row.get("item")
        or row.get("objeto")
        or ""
    ).strip()
    if not descricao:
        return None
    quantidade_raw = row.get("quantidade") or row.get("qtd") or 1
    try:
        quantidade = float(str(quantidade_raw).replace(",", "."))
    except ValueError:
        quantidade = 1
    return {
        "codigo": str(row.get("codigo") or row.get("código") or row.get("cod") or "").strip(),
        "descricao": descricao,
        "unidade": str(row.get("unidade") or row.get("und") or "").strip(),
        "quantidade": quantidade,
        "categoria": str(row.get("categoria") or "").strip(),
        "termo_busca": str(row.get("termo_busca") or row.get("termo") or descricao).strip(),
        "especificacao": str(row.get("especificacao") or row.get("especificação") or "").strip(),
        "status": "pendente",
    }


def _parse_csv_itens(content: bytes) -> list[dict[str, Any]]:
    text = content.decode("utf-8-sig")
    sample = text[:2048]
    try:
        reader = csv.DictReader(StringIO(text), dialect=csv.Sniffer().sniff(sample, delimiters=";,"))
    except csv.Error:
        reader = csv.DictReader(StringIO(text), delimiter=";")
    return [
        item
        for row in reader
        if (item := _linha_para_item({_normalizar_cabecalho(k): v for k, v in row.items()}))
    ]


def _parse_xlsx_itens(content: bytes) -> list[dict[str, Any]]:
    wb = load_workbook(BytesIO(content), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [_normalizar_cabecalho(value) for value in rows[0]]
    itens = []
    for values in rows[1:]:
        row = {headers[idx]: value for idx, value in enumerate(values) if idx < len(headers)}
        item = _linha_para_item(row)
        if item:
            itens.append(item)
    return itens


@app.post("/itens/importar-arquivo")
async def importar_itens_arquivo(
    processo_id: int = Form(...),
    arquivo: UploadFile = File(...),
) -> dict[str, Any]:
    if not get_processo(processo_id):
        raise HTTPException(status_code=404, detail="Pesquisa nao encontrada")
    content = await arquivo.read()
    nome = (arquivo.filename or "").lower()
    if nome.endswith(".csv"):
        itens = _parse_csv_itens(content)
    elif nome.endswith(".xlsx"):
        itens = _parse_xlsx_itens(content)
    else:
        raise HTTPException(status_code=400, detail="Envie arquivo .csv ou .xlsx")
    if not itens:
        raise HTTPException(status_code=400, detail="Nenhum item valido encontrado no arquivo")
    ids = create_itens(processo_id, itens)
    return {
        "ids": ids,
        "itens": list_itens(processo_id),
        "resumo_itens": resumo_itens(processo_id),
    }


@app.get("/itens/{item_id}")
def obter_item(item_id: int) -> dict[str, Any]:
    item = get_item(item_id)
    if not item:
        raise HTTPException(status_code=404, detail="Item nao encontrado")
    return {
        "item": item,
        "fontes": list_fontes_item(item_id),
        "resumo": resumo_item(item_id),
    }


@app.delete("/itens/{item_id}")
def excluir_item(item_id: int) -> dict[str, Any]:
    deleted = delete_item(item_id)
    if not deleted:
        raise HTTPException(status_code=404, detail="Item nao encontrado")
    return {"status": "deleted", "id": item_id}


@app.get("/processos/{processo_id}/checklist")
def checklist(processo_id: int) -> dict[str, Any]:
    data = export_snapshot(processo_id)
    if not data:
        raise HTTPException(status_code=404, detail="Processo nao encontrado")
    itens = gerar_checklist(data)
    progresso = itens[0]["progresso"] if itens else {"concluidos": 0, "total": 0}
    return {"checklist": itens, "progresso": progresso}


@app.get("/processos/{processo_id}/comparabilidade")
def comparabilidade(processo_id: int) -> dict[str, Any]:
    data = export_snapshot(processo_id)
    if not data:
        raise HTTPException(status_code=404, detail="Processo nao encontrado")
    return avaliar_comparabilidade(data)


@app.post("/fontes")
def criar_fonte(req: FontePrecoRequest) -> dict[str, Any]:
    if not get_processo(req.processo_id):
        raise HTTPException(status_code=404, detail="Processo nao encontrado")
    if req.item_id and not get_item(req.item_id):
        raise HTTPException(status_code=404, detail="Item nao encontrado")
    fonte_id = create_fonte(req.model_dump())
    return {
        "id": fonte_id,
        "fontes": list_fontes(req.processo_id),
        "resumo": resumo_pesquisa(req.processo_id),
        "resumo_itens": resumo_itens(req.processo_id),
    }


@app.delete("/fontes/{fonte_id}")
def excluir_fonte(fonte_id: int) -> dict[str, Any]:
    deleted = delete_fonte(fonte_id)
    if not deleted:
        raise HTTPException(status_code=404, detail="Fonte nao encontrada")
    return {"status": "deleted", "id": fonte_id}


@app.post("/atas")
def criar_ata(req: AtaRequest) -> dict[str, Any]:
    if not get_processo(req.processo_id):
        raise HTTPException(status_code=404, detail="Processo nao encontrado")
    ata_id = create_ata(req.model_dump())
    return {
        "id": ata_id,
        "atas": list_atas(req.processo_id),
        "resumo_atas": resumo_atas(req.processo_id),
    }


@app.delete("/atas/{ata_id}")
def excluir_ata(ata_id: int) -> dict[str, Any]:
    deleted = delete_ata(ata_id)
    if not deleted:
        raise HTTPException(status_code=404, detail="Ata nao encontrada")
    return {"status": "deleted", "id": ata_id}


@app.get("/processos/{processo_id}/atas")
def atas(processo_id: int) -> dict[str, Any]:
    if not get_processo(processo_id):
        raise HTTPException(status_code=404, detail="Processo nao encontrado")
    return {
        "atas": list_atas(processo_id),
        "resumo_atas": resumo_atas(processo_id),
    }


@app.post("/pncp/buscar")
def pncp_buscar(req: PncpBuscaRequest) -> dict[str, Any]:
    return buscar_contratacoes(
        termo=req.termo,
        data_inicial=req.data_inicial or None,
        data_final=req.data_final or None,
        uf=req.uf or None,
        modalidade_id=req.modalidade_id or None,
        pagina=req.pagina,
        tamanho_pagina=req.tamanho_pagina,
    )


@app.post("/pncp/buscar-contexto")
def pncp_buscar_contexto(req: PncpContextoBuscaRequest) -> dict[str, Any]:
    texto_ref = req.texto_referencia.strip()
    termo = req.termo.strip()

    if req.item_id is not None:
        item = get_item(req.item_id)
        if not item:
            raise HTTPException(status_code=404, detail="Item nao encontrado")
        partes_ref = [
            str(item.get("descricao") or "").strip(),
            str(item.get("especificacao") or "").strip(),
            str(item.get("termo_busca") or "").strip(),
        ]
        if not texto_ref:
            texto_ref = " ".join(p for p in partes_ref if p).strip()
        if not termo:
            termo = (str(item.get("termo_busca") or "").strip()) or (
                str(item.get("descricao") or "").strip()[:200]
            )

    if not termo and not texto_ref:
        raise HTTPException(
            status_code=400,
            detail="Informe termo, texto de referencia ou item_id",
        )
    if not texto_ref:
        texto_ref = termo

    return buscar_contratacoes_contextual(
        termo=termo or texto_ref,
        texto_referencia=texto_ref,
        data_inicial=req.data_inicial or None,
        data_final=req.data_final or None,
        uf=req.uf or None,
        ufs=req.ufs or None,
        modalidade_id=req.modalidade_id or None,
        tamanho_pagina=req.tamanho_pagina,
        buscar_variantes=req.buscar_variantes,
        max_consultas=req.max_consultas,
        estrategia=req.estrategia,
        similaridade_minima=req.similaridade_minima,
    )


@app.post("/pncp/rascunhos")
def criar_pncp_rascunho(req: PncpRascunhoRequest) -> dict[str, Any]:
    if not get_processo(req.processo_id):
        raise HTTPException(status_code=404, detail="Processo nao encontrado")
    if req.item_id and not get_item(req.item_id):
        raise HTTPException(status_code=404, detail="Item nao encontrado")
    rascunho_id = create_pncp_rascunho(req.model_dump())
    return {
        "id": rascunho_id,
        "rascunhos": list_pncp_rascunhos(req.processo_id),
    }


@app.get("/processos/{processo_id}/pncp-rascunhos")
def listar_pncp_rascunhos(processo_id: int, item_id: int | None = None, status: str | None = None) -> dict[str, Any]:
    if not get_processo(processo_id):
        raise HTTPException(status_code=404, detail="Processo nao encontrado")
    return {"rascunhos": list_pncp_rascunhos(processo_id, item_id=item_id, status=status)}


@app.post("/pncp/rascunhos/{rascunho_id}/fonte")
def converter_rascunho_em_fonte(rascunho_id: int, req: RascunhoFonteRequest) -> dict[str, Any]:
    rascunho = get_pncp_rascunho(rascunho_id)
    if not rascunho:
        raise HTTPException(status_code=404, detail="Rascunho nao encontrado")
    fonte_id = create_fonte_from_pncp_rascunho(rascunho_id, req.model_dump())
    if fonte_id is None:
        raise HTTPException(status_code=404, detail="Rascunho nao encontrado")
    processo_id = int(rascunho["processo_id"])
    return {
        "id": fonte_id,
        "fontes": list_fontes(processo_id),
        "resumo": resumo_pesquisa(processo_id),
        "resumo_itens": resumo_itens(processo_id),
        "rascunhos": list_pncp_rascunhos(processo_id),
    }


@app.delete("/pncp/rascunhos/{rascunho_id}")
def excluir_pncp_rascunho(rascunho_id: int) -> dict[str, Any]:
    deleted = delete_pncp_rascunho(rascunho_id)
    if not deleted:
        raise HTTPException(status_code=404, detail="Rascunho nao encontrado")
    return {"status": "deleted", "id": rascunho_id}


@app.get("/processos/{processo_id}/resumo")
def resumo(processo_id: int) -> dict[str, Any]:
    if not get_processo(processo_id):
        raise HTTPException(status_code=404, detail="Processo nao encontrado")
    return resumo_pesquisa(processo_id)


@app.get("/processos/{processo_id}/revisao")
def revisao_final(processo_id: int) -> dict[str, Any]:
    data = export_snapshot(processo_id)
    if not data:
        raise HTTPException(status_code=404, detail="Processo nao encontrado")
    return gerar_revisao_itens(data)


@app.get("/processos/{processo_id}/snapshot")
def snapshot(processo_id: int) -> dict[str, Any]:
    data = export_snapshot(processo_id)
    if not data:
        raise HTTPException(status_code=404, detail="Processo nao encontrado")
    return data


@app.get("/processos/{processo_id}/export/xlsx")
def exportar_xlsx(processo_id: int) -> FileResponse:
    path = gerar_xlsx_processo(processo_id)
    if not path:
        raise HTTPException(status_code=404, detail="Processo nao encontrado")
    return FileResponse(
        path,
        filename=path.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.get("/processos/{processo_id}/export/md")
def exportar_relatorio_md(processo_id: int) -> FileResponse:
    path = gerar_relatorio_markdown(processo_id)
    if not path:
        raise HTTPException(status_code=404, detail="Processo nao encontrado")
    return FileResponse(
        path,
        filename=path.name,
        media_type="text/markdown; charset=utf-8",
    )


@app.get("/processos/{processo_id}/export/html")
def exportar_relatorio_html(processo_id: int) -> FileResponse:
    path = gerar_relatorio_html(processo_id)
    if not path:
        raise HTTPException(status_code=404, detail="Processo nao encontrado")
    return FileResponse(path, filename=path.name, media_type="text/html; charset=utf-8")


@app.get("/processos/{processo_id}/export/docx")
def exportar_relatorio_docx(processo_id: int) -> FileResponse:
    path = gerar_relatorio_docx(processo_id)
    if not path:
        raise HTTPException(status_code=404, detail="Processo nao encontrado")
    return FileResponse(
        path,
        filename=path.name,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
